#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""根据 project.json 与 test_case.csv 生成 Excel 测试用例。
表头：项目名称、项目分级、体测项目｜体能项目、操作步骤、算法预期结果、实际结果。
项目列表以 project.json 的 projectList 为准；项目分级来自 test_case.csv（等级）；exercise 0=体测项目，1=体能项目。
"""
import csv
import json
import os

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
except ImportError:
    print("请先安装: pip install openpyxl")
    raise

# 统一操作步骤模板（{} 填入项目名称，展示为【项目名称】）
STEPS_TEMPLATE = (
    "1. 切换并执行【{}】项目\n"
    "2. 在规定的时间内完成项目\n"
    "3. 检查计数情况"
)

# API 项目名称 -> test_case.csv 项目名称，用于匹配等级（CSV 与 API 命名不一致时）
DESC_TO_CSV_NAME = {
    "跳绳": "双脚跳绳",
    "篮球绕杆往返": "篮球绕杆",
    "50x8往返跑": "50米×8往返跑",
}


def load_project_list(project_json_path):
    """从 project.json 读取 projectList，返回 [(desc, exercise), ...]。exercise 缺省为 0。"""
    with open(project_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    pl = data.get("data", {}).get("projectList", [])
    out = []
    for p in pl:
        desc = (p.get("desc") or "").strip()
        if not desc:
            continue
        ex = p.get("exercise", 0)
        out.append((desc, 1 if ex == 1 else 0))
    return out


def load_level_from_csv(csv_path):
    """从 test_case.csv 读取 项目名称 -> 等级，用于项目分级。"""
    level_map = {}
    if not os.path.isfile(csv_path):
        return level_map
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            name = (row.get("项目名称") or "").strip()
            level = (row.get("等级") or "").strip()
            if name:
                level_map[name] = level or "—"
    return level_map


def algorithm_expected(project_name):
    """根据项目名称返回算法预期结果，分行编号展示（1. 2. 3.）。"""
    if "跳绳" in project_name:
        return "1. 算法正确识别跳绳次数与状态；\n2. 计数准确，漏检率在允许范围内；\n3. 成绩与规则一致。"
    if "跑" in project_name or "米" in project_name:
        return "1. 算法正确识别跑道与圈数；\n2. 漏圈/多圈判断正确；\n3. 成绩与规则一致。"
    if "跳远" in project_name:
        return "1. 起跳线、落点识别准确；\n2. 距离计算正确；\n3. 成绩有效。"
    if "跳" in project_name:
        return "1. 起跳/落点或动作识别准确；\n2. 计数或距离计算正确；\n3. 成绩有效。"
    if "引体" in project_name or "仰卧" in project_name or "俯卧" in project_name:
        return "1. 动作识别与计数准确；\n2. 无效动作过滤正确；\n3. 成绩符合规则。"
    if "排球" in project_name or "篮球" in project_name or "足球" in project_name or "实心球" in project_name:
        return "1. 动作与轨迹识别正确；\n2. 计数/计时与规则一致；\n3. 无漏检误检。"
    if "坐位体前屈" in project_name or "肺活量" in project_name:
        return "1. 读数/测量结果识别正确；\n2. 单位与规则一致；\n3. 结果可复现。"
    if "身高" in project_name or "体重" in project_name:
        return "1. 数值识别正确；\n2. 单位正确；\n3. 与设备/规则一致。"
    return "1. 算法正确识别该项目动作/轨迹；\n2. 人数、成绩等统计与规则一致；\n3. 日志无异常，结果可复现。"


def main():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_json_path = os.path.join(script_dir, "project.json")
    csv_path = os.path.join(script_dir, "test_case.csv")
    out_path = os.path.join(script_dir, "test_case_用例.xlsx")

    if not os.path.isfile(project_json_path):
        print(f"未找到: {project_json_path}")
        return

    projects = load_project_list(project_json_path)
    level_map = load_level_from_csv(csv_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例"

    # 表头：项目名称、项目分级、体测项目｜体能项目、操作步骤、算法预期结果、实际结果
    headers = ["项目名称", "项目分级", "体测项目｜体能项目", "操作步骤", "算法预期结果", "实际结果"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (desc, exercise) in enumerate(projects, start=2):
        category = "体能项目" if exercise == 1 else "体测项目"
        # 体能项目统一为 P2；体测项目按 test_case.csv 匹配等级
        if exercise == 1:
            level = "P2"
        else:
            csv_name = DESC_TO_CSV_NAME.get(desc, desc)
            level = level_map.get(csv_name) or level_map.get(desc) or "—"
        steps = STEPS_TEMPLATE.format(desc)
        expected = algorithm_expected(desc)

        ws.cell(row=i, column=1, value=desc)
        ws.cell(row=i, column=2, value=level)
        ws.cell(row=i, column=3, value=category)
        ws.cell(row=i, column=4, value=steps)
        ws.cell(row=i, column=5, value=expected)
        ws.cell(row=i, column=6, value="")

        for c in range(1, 7):
            ws.cell(row=i, column=c).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=i, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["E"].width = 48
    ws.column_dimensions["F"].width = 36

    wb.save(out_path)
    print(f"已生成: {out_path}（共 {len(projects)} 条）")


if __name__ == "__main__":
    main()
